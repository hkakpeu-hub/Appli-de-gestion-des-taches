"""Exportation Excel (.xlsx).

Le classeur produit contient deux feuilles :
 - « actions journalières » : même disposition que le fichier de travail d'origine (en-têtes en ligne 4,
   colonnes B à I, statuts « 1-Terminé / 2-En cours / 3-Non dem ») -> directement utilisable à la place de l'ancien fichier ;
 - « Toutes les données » : toutes les informations (priorité, horaires, durées, responsable, commentaires…).
Les deux feuilles sont ré-importables par l'application.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..core import rules
from ..core.models import KIND_DONE, KIND_IN_PROGRESS, KIND_TODO, PRIORITY_LABELS, Task
from ..core.tasks import TaskService

LEGACY_STATUS = {KIND_DONE: "1-Terminé", KIND_IN_PROGRESS: "2-En cours", KIND_TODO: "3-Non dem"}
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D0D7E2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DATE_FMT = "dd/mm/yyyy"


def _d(d):
    return datetime.combine(d, datetime.min.time()) if d else None


def _comments_text(svc: TaskService, t: Task) -> str:
    items = svc.comments(t.id)
    return "\n".join(f"[{c.created_at:%d/%m/%Y}] {c.body}" for c in items)


def _style_header(ws, row: int, first_col: int, last_col: int) -> None:
    for c in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def export_tasks(svc: TaskService, tasks: Iterable[Task], path: Path, today=None) -> int:
    """Écrit ``tasks`` dans ``path``. Retourne le nombre de lignes exportées."""
    tasks = list(tasks)
    today = today or svc.clock().date()
    wb = Workbook()

    # --- feuille compatible avec l'ancien fichier ----------------------------------------------
    ws = wb.active
    ws.title = "actions journalières"
    ws["B1"] = "ACTIONS AU QUOTIDIEN -- Export de l'application Suivi des tâches"
    ws["B1"].font = Font(bold=True, size=14)
    ws["B2"] = f"Exporté le {svc.clock():%d/%m/%Y à %H:%M} · {len(tasks)} tâche(s)"
    heads = ["Dates", "Action", "Dead line", "Projet", "statut", "Description",
             "Commetaire(Difficultés, apprentissage)", "Date complémentaire"]
    for i, h in enumerate(heads):
        ws.cell(row=4, column=2 + i, value=h)
    _style_header(ws, 4, 2, 9)
    ordered = sorted(tasks, key=lambda t: (t.start_date or t.created_at.date(), t.id))
    for r, t in enumerate(ordered, start=5):
        vals = [
            _d(t.start_date or t.created_at.date()), t.title, _d(t.deadline), t.project_name,
            LEGACY_STATUS.get(t.status_kind, t.status_label), t.description, _comments_text(svc, t), _d(t.extra_date),
        ]
        for i, v in enumerate(vals):
            cell = ws.cell(row=r, column=2 + i, value=v)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=i in (1, 5, 6))
            if i in (0, 2, 7):
                cell.number_format = DATE_FMT
    for col, w in zip("BCDEFGHI", (13, 60, 13, 16, 14, 50, 50, 18)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"B4:I{max(4, 4 + len(ordered))}"

    # --- feuille complète -----------------------------------------------------------------------
    full = wb.create_sheet("Toutes les données")
    cols = ["ID", "Intitulé", "Projet", "Description", "Responsable", "Statut", "Priorité", "Date de création",
            "Date de début", "Heure de début", "Heure de fin", "Deadline", "Date de réalisation", "Durée estimée (min)",
            "Durée réelle (min)", "Date complémentaire", "Jours de retard", "Dernière modification", "Commentaires"]
    full.append(cols)
    _style_header(full, 1, 1, len(cols))
    for t in tasks:
        full.append([
            t.id, t.title, t.project_name, t.description, t.assignee, t.status_label, PRIORITY_LABELS.get(t.priority),
            t.created_at, _d(t.start_date), t.start_time.strftime("%H:%M") if t.start_time else None,
            t.end_time.strftime("%H:%M") if t.end_time else None, _d(t.deadline), t.completed_at,
            t.est_minutes, t.actual_minutes, _d(t.extra_date), rules.overdue_days(t, today) or None,
            t.updated_at, _comments_text(svc, t),
        ])
    date_cols = {8, 9, 12, 13, 16}
    by_id = {t.id: t for t in tasks}
    for row in full.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (2, 4, 19))
            if cell.column in date_cols:
                cell.number_format = DATE_FMT
            elif cell.column == 18:
                cell.number_format = "dd/mm/yyyy hh:mm"
        st = by_id.get(row[0].value)
        if st:
            row[5].fill = PatternFill("solid", fgColor=st.status_color.lstrip("#").upper())
            row[5].font = Font(color="FFFFFF", bold=True)
    for i, w in enumerate((7, 60, 16, 50, 16, 14, 11, 16, 13, 11, 11, 13, 16, 12, 12, 16, 10, 18, 50), start=1):
        full.column_dimensions[get_column_letter(i)].width = w
    full.freeze_panes = "C2"
    full.auto_filter.ref = full.dimensions

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)  # écriture dans un fichier temporaire : pas de fichier tronqué en cas d'erreur
    tmp.replace(path)
    return len(tasks)
