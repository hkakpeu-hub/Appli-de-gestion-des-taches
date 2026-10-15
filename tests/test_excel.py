import os
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

from suivi.excel.exporter import export_tasks
from suivi.excel.importer import ExcelImporter, ImportFailed, parse_date, parse_time


def make_legacy(path: Path):
    """Reproduit la structure du fichier de travail d'origine."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "actions journalières"
    ws["B1"] = "ACTIONS AU QUOTIDIEN"
    ws.append([])
    ws.append([])
    ws.append([None, "Dates", "Action", "Dead line", "Projet", "statut", "Description",
               "Commetaire(Difficultés, apprentissage)", None])
    ws.append([1, datetime(2025, 1, 27), "Restitution du rapport", datetime(2025, 1, 27), "TNP", "1-Terminé", None, None, None])
    ws.append([2, datetime(2025, 1, 27), "Participer au COSTRAT", datetime(2025, 1, 27), "Centaures", "2-En cours",
               "Stratégie de recette", "Attente de validation", datetime(2025, 2, 17)])
    ws.append([3, datetime(2025, 1, 28), "Rédiger le CR", datetime(2025, 2, 6), "Centaures", "3-Non dem", None, None, None])
    ws.append([4, "pas une date", "Ligne à date invalide", "31/02/2025", "TNP", "Statut bizarre", None, None, None])
    ws.append([5, datetime(2025, 1, 29), None, None, "TNP", None, "sans titre", None, None])  # rejetée
    ws.append([6, datetime(2025, 1, 27), "Restitution du rapport", datetime(2025, 1, 27), "TNP", "1-Terminé", None, None, None])  # doublon fichier
    perso = wb.create_sheet("Actions personnelles")
    perso.append(["ACTIONS A FAIRE"])
    perso.append([])
    perso.append([None, "Dates", "Action", "Heure de début", "Heure de fin", "statut", "Description"])
    perso.append([None, None, "Créer un compte bancaire", datetime(2025, 10, 30), None, None, "le mois prochain"])
    other = wb.create_sheet("Compétences")
    other.append(["ID", "Date", "Compétence", "Statut"])
    other.append([1, None, "Stratégie de migration", None])
    wb.save(path)


@pytest.fixture
def legacy(tmp_path):
    p = tmp_path / "legacy.xlsx"
    make_legacy(p)
    return p


def test_parsers():
    assert parse_date("27/01/2025")[0] == date(2025, 1, 27)
    assert parse_date("2025-01-27")[0] == date(2025, 1, 27)
    assert parse_date("31/02/2025")[0] is None and parse_date("#REF!") == (None, None)
    assert parse_time("9h30")[0].hour == 9 and parse_time(0.5)[0].hour == 12
    assert parse_time(datetime(2025, 1, 1))[1] == date(2025, 1, 1)  # date dans colonne d'heure


def test_preview_detects_sheets_and_problems(ctx, legacy):
    pv = ctx.importer.preview(legacy)
    sel = {s.name: s.selected for s in pv.sheets}
    assert sel == {"actions journalières": True, "Actions personnelles": True, "Compétences": False}
    c = pv.counts()
    assert c["new"] == 5 and c["invalid"] == 1 and c["duplicate"] == 1
    bad = next(r for r in pv.rows if r.data and r.data.title == "Ligne à date invalide")
    assert any("date invalide" in m for m in bad.messages) and any("statut inconnu" in m for m in bad.messages)
    # la colonne I sans en-tête est reconnue comme date complémentaire
    assert pv.sheets[0].columns.get("extra_date") == 8
    # une date dans la colonne « Heure de début » devient une deadline
    perso = next(r for r in pv.rows if r.sheet == "Actions personnelles")
    assert perso.data.deadline == date(2025, 10, 30) and perso.data.project_name == "Personnel"
    assert ctx.tasks.all() == []  # la prévisualisation n'écrit rien


def test_import_maps_statuses_and_keeps_data(ctx, legacy):
    rep = ctx.importer.commit(ctx.importer.preview(legacy))
    assert (rep.imported, rep.rejected, rep.ignored) == (5, 1, 1)
    assert rep.backup is not None and rep.backup.exists()
    tasks = {t.title: t for t in ctx.tasks.all()}
    assert tasks["Restitution du rapport"].status_kind == "done"
    cos = tasks["Participer au COSTRAT"]
    assert cos.status_kind == "in_progress" and cos.extra_date == date(2025, 2, 17)
    assert [c.body for c in ctx.tasks.comments(cos.id)] == ["Attente de validation"]
    assert tasks["Rédiger le CR"].status_kind == "todo"
    assert ctx.backups.list(), "une sauvegarde est faite avant l'import"


def test_reimport_creates_no_duplicates(ctx, legacy):
    ctx.importer.commit(ctx.importer.preview(legacy))
    n = len(ctx.tasks.all())
    pv = ctx.importer.preview(legacy)
    assert pv.counts()["new"] == 0
    rep = ctx.importer.commit(pv)
    assert rep.imported == 0 and len(ctx.tasks.all()) == n


def test_update_existing_option(ctx, legacy):
    ctx.importer.commit(ctx.importer.preview(legacy))
    wb = openpyxl.load_workbook(legacy)
    wb["actions journalières"]["F7"] = "1-Terminé"  # « Rédiger le CR » passe à Terminé
    wb.save(legacy)
    assert ctx.importer.preview(legacy).counts()["update"] == 0
    pv = ctx.importer.preview(legacy, update_existing=True)
    assert pv.counts()["update"] == 1
    rep = ctx.importer.commit(pv)
    assert rep.updated == 1
    assert next(t for t in ctx.tasks.all() if t.title == "Rédiger le CR").status_kind == "done"


def test_failed_import_is_atomic(ctx, legacy, monkeypatch):
    pv = ctx.importer.preview(legacy)
    original = ctx.tasks.create
    calls = {"n": 0}

    def flaky(*a, **k):
        calls["n"] += 1
        if calls["n"] == 3:
            raise RuntimeError("panne simulée")
        return original(*a, **k)

    monkeypatch.setattr(ctx.tasks, "create", flaky)
    with pytest.raises(ImportFailed):
        ctx.importer.commit(pv)
    assert ctx.tasks.all() == []  # rien n'a été importé à moitié


def test_export_roundtrip_and_filtered(ctx, legacy, tmp_path):
    ctx.importer.commit(ctx.importer.preview(legacy))
    tasks = ctx.tasks.all()
    out = tmp_path / "export.xlsx"
    assert export_tasks(ctx.tasks, tasks, out) == len(tasks)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["actions journalières", "Toutes les données"]
    ws = wb["actions journalières"]
    assert [c.value for c in ws[4]][1:4] == ["Dates", "Action", "Dead line"]
    statuses = {ws.cell(r, 6).value for r in range(5, 5 + len(tasks))}
    assert {"1-Terminé", "2-En cours", "3-Non dem"} <= statuses
    # ré-import de l'export dans une base vierge : mêmes tâches, sans doublon
    from suivi.context import AppContext
    c2 = AppContext(tmp_path / "other.db", tmp_path / "bk2", ctx.clock)
    pv = c2.importer.preview(out)
    assert [s.name for s in pv.sheets if s.selected] == ["Toutes les données"]
    c2.importer.commit(pv)
    assert sorted(t.title for t in c2.tasks.all()) == sorted(t.title for t in tasks)
    assert c2.importer.preview(out).counts()["new"] == 0
    c2.close()
    # export filtré
    from suivi.core.tasks import TaskFilter
    only = ctx.tasks.list(TaskFilter(text="costrat"))
    assert export_tasks(ctx.tasks, only, tmp_path / "f.xlsx") == 1


def test_unreadable_file(ctx, tmp_path):
    bad = tmp_path / "x.xlsx"
    bad.write_text("pas un excel")
    with pytest.raises(ValueError):
        ctx.importer.preview(bad)
    with pytest.raises(ValueError):
        ctx.importer.preview(tmp_path / "x.csv")


REAL = Path(os.environ.get("SUIVI_REAL_XLSX", r"C:\Users\HermannKAKPEU\Downloads\Matrice des compétences suivi des actions.xlsx"))


@pytest.mark.skipif(not REAL.exists(), reason="fichier réel absent")
def test_real_workbook(ctx):
    pv = ctx.importer.preview(REAL)
    assert pv.counts()["invalid"] == 0 and pv.counts()["new"] > 400
    ctx.importer.commit(pv)
    assert ctx.importer.preview(REAL).counts()["new"] == 0
